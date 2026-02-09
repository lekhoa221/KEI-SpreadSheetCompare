from PyQt6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QTableWidget,
    QTableWidgetItem,
    QLabel,
    QHBoxLayout,
    QPushButton,
    QTableWidgetSelectionRange,
    QLineEdit,
    QTextEdit,
    QMenu,
    QMessageBox,
    QToolButton,
    QStyle,
    QSlider,
    QButtonGroup,
    QFrame,
    QComboBox,
    QProgressBar,
    QSplitter,
)
from PyQt6.QtGui import QColor, QIcon, QFont
from PyQt6.QtCore import Qt, pyqtSignal, QEvent
from ui.excel.delegates import BorderDelegate
from PyQt6.QtGui import QGuiApplication
import numbers
import math
import re
import os
import subprocess
from pathlib import Path
from difflib import SequenceMatcher
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
import html
from PyQt6.QtGui import QAction, QKeySequence
from ui.excel.find_dialog import FindDialog
from ui.excel.filter_header import FilterHeader


class ResultView(QWidget):
    swapRequested = pyqtSignal()
    viewModeChanged = pyqtSignal(str)
    sheetChangeRequested = pyqtSignal(int, str)
    def __init__(self, on_back):
        super().__init__()
        # ... (layout setup)
        self.layout = QVBoxLayout(self)
        self.layout.setContentsMargins(12, 4, 12, 6)
        self.layout.setSpacing(4)

        self.decimal_overrides1 = {}
        self.decimal_overrides2 = {}
        self.show_differences = False
        self.show_diff_content = True
        self.show_diff_format = True
        self.show_diff_formula = False
        self.ignore_case = False
        self.ignore_whitespace = False
        self.ignore_number_format = False
        self.only_changes = False
        self.diff_map = {}
        self.diff_rows_content = set()
        self.diff_rows_format = set()
        self.diff_rows_formula = set()
        self.df1 = None
        self.df2 = None
        self.styles1 = None
        self.styles2 = None
        self.merges1 = None
        self.merges2 = None
        self.formulas1 = {}
        self.formulas2 = {}
        self.change_set = None
        self.total_rows = 0
        self.total_cols = 0
        self.active_grid = None
        self.active_grid = None
        self.highlighted_refs = []
        self.highlighted_grid = None
        self.ref_colors = ["#ef4444", "#3b82f6", "#10b981", "#f59e0b", "#8b5cf6", "#ec4899", "#06b6d4"]
        self.sheet1_current = None
        self.sheet2_current = None
        self.file1_path = None
        self.file2_path = None
        
        # Formula Bar
        formula_layout = QHBoxLayout()
        formula_layout.setContentsMargins(0, 0, 0, 0)
        formula_layout.setSpacing(4)

        self.name_box = QLineEdit()
        self.name_box.setObjectName("NameBox")
        self.name_box.setReadOnly(True)
        self.name_box.setPlaceholderText("Name Box")
        self.name_box.setFixedWidth(120)

        self.formula_bar = QTextEdit()
        self.formula_bar.setObjectName("FormulaBar")
        self.formula_bar.setReadOnly(True)
        self.formula_bar.setPlaceholderText("Formula")
        self.formula_bar.setFixedHeight(28) # Approximate height for single line
        self.formula_bar.setLineWrapMode(QTextEdit.LineWrapMode.NoWrap)
        self.formula_bar.setStyleSheet("QTextEdit { padding: 4px; }")

        # Sync Toggle
        self.sync_toggle = QToolButton()
        self.sync_toggle.setText("Sync")
        self.sync_toggle.setCheckable(True)
        self.sync_toggle.setChecked(False) # Default False as requested
        self.sync_toggle.setToolTip("Synchronize scrolling and resizing between sheets")
        # Apply style directly for visibility or use qss
        self.sync_toggle.setStyleSheet("""
            QToolButton { font-size: 11px; font-weight: bold; border: 1px solid #ccc; border-radius: 4px; padding: 2px 8px; }
            QToolButton:checked { background-color: #107c41; color: white; border-color: #107c41; }
        """)
        
        # Sync Options State
        self.sync_scroll_enabled = True
        self.sync_row_height_enabled = True
        self.sync_col_width_enabled = True

        # Custom Context Menu for Sync Button
        self.sync_toggle.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.sync_toggle.customContextMenuRequested.connect(self.show_sync_params_menu)

        formula_layout.addWidget(self.name_box)
        formula_layout.addWidget(self.formula_bar, 1)
        
        # Swap Button in Result View
        self.swap_btn = QToolButton()
        self.swap_btn.setIcon(QIcon(os.path.join(os.path.dirname(__file__), "..", "assets", "icons", "diff-toggle.svg"))) # Placeholder or reuse reload icon
        # Actually I should allow passing icon or load it. 
        # For now let's just use text or standard icon if I can't easily access the asset loader here.
        self.swap_btn.setText("Swap")
        self.swap_btn.setToolTip("Swap Sides and Re-compare")
        self.swap_btn.clicked.connect(self.request_swap)
        self.swap_btn.setStyleSheet("""
            QToolButton { border: 1px solid #e2e8f0; border-radius: 4px; padding: 2px 8px; font-size: 11px; font-weight: bold; background: #fff; }
            QToolButton:hover { background-color: #f1f5f9; }
        """)
        formula_layout.addWidget(self.swap_btn)

        formula_layout.addWidget(self.sync_toggle)
        self.layout.addLayout(formula_layout)
        
        # Split View for Side-by-Side
        self.splitter = QSplitter(Qt.Orientation.Horizontal)
        self.splitter.setChildrenCollapsible(False)
        self.splitter.setHandleWidth(6)

        # Grid 1
        self.grid1_frame = QFrame()
        self.grid1_frame.setObjectName("ResultSide")
        self.grid1_container = QVBoxLayout(self.grid1_frame)
        self.grid1_container.setSpacing(2)
        self.grid1_container.setContentsMargins(8, 6, 8, 8)

        self.grid1_header = QHBoxLayout()
        self.grid1_header.setContentsMargins(0, 0, 0, 0)
        self.grid1_header.setSpacing(6)

        self.grid1_label = QLabel("Original")
        self.grid1_label.setProperty("panelTitle", True)
        self.grid1_label.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.grid1_label.customContextMenuRequested.connect(lambda pos: self.show_file_context_menu(1, pos))
        self.grid1_sheet_btn = QToolButton()
        self.grid1_sheet_btn.setObjectName("SheetButton")
        self.grid1_sheet_btn.setAutoRaise(True)
        self.grid1_sheet_btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextOnly)
        self.grid1_sheet_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.grid1_sheet_btn.clicked.connect(lambda: self.show_sheet_combo(1))

        self.grid1_sheet_combo = QComboBox()
        self.grid1_sheet_combo.setObjectName("SheetCombo")
        self.grid1_sheet_combo.setEnabled(False)
        self.grid1_sheet_combo.setMinimumWidth(160)
        self.grid1_sheet_combo.setVisible(False)
        self.grid1_sheet_combo.activated.connect(lambda idx: self.on_sheet_combo_selected(1, idx))
        self.grid1_sheet_combo.installEventFilter(self)

        self.grid1 = QTableWidget()
        self.grid1.setItemDelegate(BorderDelegate(self.grid1))
        self.grid1.setWordWrap(False)
        self.grid1.verticalHeader().setDefaultAlignment(Qt.AlignmentFlag.AlignCenter)
        self.grid1.verticalHeader().setFixedWidth(40)

        self.grid1_header.addWidget(self.grid1_label)
        self.grid1_header.addWidget(self.grid1_sheet_btn)
        self.grid1_header.addWidget(self.grid1_sheet_combo)
        self.grid1_header.addStretch()
        self.grid1_container.addLayout(self.grid1_header)
        self.grid1_container.addWidget(self.grid1)

        # Grid 2
        self.grid2_frame = QFrame()
        self.grid2_frame.setObjectName("ResultSide")
        self.grid2_container = QVBoxLayout(self.grid2_frame)
        self.grid2_container.setSpacing(2)
        self.grid2_container.setContentsMargins(8, 6, 8, 8)

        self.grid2_header = QHBoxLayout()
        self.grid2_header.setContentsMargins(0, 0, 0, 0)
        self.grid2_header.setSpacing(6)

        self.grid2_label = QLabel("Modified")
        self.grid2_label.setProperty("panelTitle", True)
        self.grid2_label.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.grid2_label.customContextMenuRequested.connect(lambda pos: self.show_file_context_menu(2, pos))
        self.grid2_sheet_btn = QToolButton()
        self.grid2_sheet_btn.setObjectName("SheetButton")
        self.grid2_sheet_btn.setAutoRaise(True)
        self.grid2_sheet_btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextOnly)
        self.grid2_sheet_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.grid2_sheet_btn.clicked.connect(lambda: self.show_sheet_combo(2))

        self.grid2_sheet_combo = QComboBox()
        self.grid2_sheet_combo.setObjectName("SheetCombo")
        self.grid2_sheet_combo.setEnabled(False)
        self.grid2_sheet_combo.setMinimumWidth(160)
        self.grid2_sheet_combo.setVisible(False)
        self.grid2_sheet_combo.activated.connect(lambda idx: self.on_sheet_combo_selected(2, idx))
        self.grid2_sheet_combo.installEventFilter(self)

        self.grid2 = QTableWidget()
        self.grid2.setItemDelegate(BorderDelegate(self.grid2))
        self.grid2.setWordWrap(False)
        self.grid2.verticalHeader().setDefaultAlignment(Qt.AlignmentFlag.AlignCenter)
        self.grid2.verticalHeader().setFixedWidth(40)

        self.grid2_header.addWidget(self.grid2_label)
        self.grid2_header.addWidget(self.grid2_sheet_btn)
        self.grid2_header.addWidget(self.grid2_sheet_combo)
        self.grid2_header.addStretch()
        self.grid2_container.addLayout(self.grid2_header)
        self.grid2_container.addWidget(self.grid2)

        self.splitter.addWidget(self.grid1_frame)
        self.splitter.addWidget(self.grid2_frame)
        self.splitter.setStretchFactor(0, 1)
        self.splitter.setStretchFactor(1, 1)
        
        self.layout.addWidget(self.splitter, 1)

        # Status Bar (View Mode + Zoom)
        self.status_bar = QFrame()
        self.status_bar.setObjectName("ResultStatusBar")
        self.status_bar.setFixedHeight(28)  # Make it thinner
        status_layout = QHBoxLayout(self.status_bar)
        status_layout.setContentsMargins(10, 0, 10, 0) # Minimal vertical margins
        status_layout.setSpacing(8)

        self.status_info = QLabel("Ready")
        self.status_info.setObjectName("StatusInfo")
        status_layout.addWidget(self.status_info)
        status_layout.addStretch()

        # View mode buttons
        self.view_group = QButtonGroup(self)
        self.view_group.setExclusive(True)

        icon_dir = os.path.join(os.path.dirname(__file__), "..", "assets", "icons")

        self.view_normal_btn = QToolButton()
        self.view_normal_btn.setIcon(QIcon(os.path.join(icon_dir, "view-normal.svg")))
        self.view_normal_btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonIconOnly)
        self.view_normal_btn.setCheckable(True)
        self.view_normal_btn.setChecked(True)
        self.view_normal_btn.setAutoRaise(True)
        self.view_normal_btn.setFixedSize(22, 22)
        self.view_normal_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.view_normal_btn.setObjectName("StatusToggle")
        self.view_normal_btn.setToolTip("Normal")
        self.view_group.addButton(self.view_normal_btn, 0)

        self.view_page_btn = QToolButton()
        self.view_page_btn.setIcon(QIcon(os.path.join(icon_dir, "view-page-break.svg")))
        self.view_page_btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonIconOnly)
        self.view_page_btn.setCheckable(True)
        self.view_page_btn.setAutoRaise(True)
        self.view_page_btn.setFixedSize(22, 22)
        self.view_page_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.view_page_btn.setObjectName("StatusToggle")
        self.view_page_btn.setToolTip("Page Break Preview")
        self.view_group.addButton(self.view_page_btn, 1)

        status_layout.addWidget(self.view_normal_btn)
        status_layout.addWidget(self.view_page_btn)

        # Loading bar (placeholder for future behavior)
        self.status_progress = QProgressBar()
        self.status_progress.setObjectName("StatusProgress")
        self.status_progress.setTextVisible(False)
        self.status_progress.setRange(0, 100)
        self.status_progress.setValue(0)
        self.status_progress.setFixedWidth(90)
        self.status_progress.setVisible(False)
        status_layout.addWidget(self.status_progress)
        self.view_mode = "normal"

        sep = QFrame()
        sep.setFrameShape(QFrame.Shape.VLine)
        sep.setFrameShadow(QFrame.Shadow.Sunken)
        sep.setFixedHeight(14)
        sep.setObjectName("StatusSeparator")
        status_layout.addWidget(sep)

        # Zoom controls
        self.zoom_out_btn = QToolButton()
        self.zoom_out_btn.setIcon(QIcon(os.path.join(icon_dir, "zoom-out.svg")))
        self.zoom_out_btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonIconOnly)
        self.zoom_out_btn.setAutoRaise(True)
        self.zoom_out_btn.setFixedSize(22, 22)
        self.zoom_out_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.zoom_out_btn.setObjectName("StatusToggle")
        self.zoom_out_btn.setToolTip("Zoom Out")
        status_layout.addWidget(self.zoom_out_btn)

        self.zoom_slider = QSlider(Qt.Orientation.Horizontal)
        self.zoom_slider.setObjectName("StatusZoomSlider")
        self.zoom_slider.setRange(25, 200)
        self.zoom_slider.setSingleStep(5)
        self.zoom_slider.setPageStep(10)
        self.zoom_slider.setValue(100)
        self.zoom_slider.setFixedWidth(120)
        self.zoom_slider.setCursor(Qt.CursorShape.PointingHandCursor)
        status_layout.addWidget(self.zoom_slider)

        self.zoom_in_btn = QToolButton()
        self.zoom_in_btn.setIcon(QIcon(os.path.join(icon_dir, "zoom-in.svg")))
        self.zoom_in_btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonIconOnly)
        self.zoom_in_btn.setAutoRaise(True)
        self.zoom_in_btn.setFixedSize(22, 22)
        self.zoom_in_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.zoom_in_btn.setObjectName("StatusToggle")
        self.zoom_in_btn.setToolTip("Zoom In")
        status_layout.addWidget(self.zoom_in_btn)

        self.zoom_label = QLabel("100%")
        self.zoom_label.setObjectName("StatusZoomLabel")
        status_layout.addWidget(self.zoom_label)

        self.layout.addWidget(self.status_bar)

        # Zoom state
        self.zoom_percent = 100
        self.zoom_percent1 = 100
        self.zoom_percent2 = 100
        self.base_table_font = None
        
        # Sync Scrolling and Selection (Same as before)
        # Sync Scrolling and Selection (Now conditional or always? User said "Sync" toggles resizing. 
        # Usually scroll sync is also desirable but user specifically mentioned height/width sync.
        # Let's assume Sync toggle affects resizing AND scrolling for a true "Independent" vs "Synced" mode.
        
        # Scroll Sync Logic utilizing the toggle
        self.grid1.verticalScrollBar().valueChanged.connect(lambda v: self.sync_scroll(self.grid1, self.grid2, Qt.Orientation.Vertical, v))
        self.grid2.verticalScrollBar().valueChanged.connect(lambda v: self.sync_scroll(self.grid2, self.grid1, Qt.Orientation.Vertical, v))
        self.grid1.horizontalScrollBar().valueChanged.connect(lambda v: self.sync_scroll(self.grid1, self.grid2, Qt.Orientation.Horizontal, v))
        self.grid2.horizontalScrollBar().valueChanged.connect(lambda v: self.sync_scroll(self.grid2, self.grid1, Qt.Orientation.Horizontal, v))

        self.is_syncing_selection = False
        self.grid1.itemSelectionChanged.connect(lambda: self.on_grid_selection_changed(self.grid1, self.grid2))
        self.grid2.itemSelectionChanged.connect(lambda: self.on_grid_selection_changed(self.grid2, self.grid1))

        # Sync Resizing
        self.is_syncing_resize = False
        self.grid1.horizontalHeader().sectionResized.connect(lambda idx, old, new: self.sync_col_resize(self.grid1, self.grid2, idx, new))
        self.grid2.horizontalHeader().sectionResized.connect(lambda idx, old, new: self.sync_col_resize(self.grid2, self.grid1, idx, new))
        self.grid1.verticalHeader().sectionResized.connect(lambda idx, old, new: self.sync_row_resize(self.grid1, self.grid2, idx, new))
        self.grid2.verticalHeader().sectionResized.connect(lambda idx, old, new: self.sync_row_resize(self.grid2, self.grid1, idx, new))


        # Context Menu
        self.grid1.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.grid1.customContextMenuRequested.connect(lambda pos: self.show_context_menu(pos, self.grid1))
        self.grid2.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.grid2.customContextMenuRequested.connect(lambda pos: self.show_context_menu(pos, self.grid2))

        # Find Dialog
        self.find_dialog = None
        
        # Shortcuts
        find_action = QAction(self)
        find_action.setShortcut(QKeySequence.StandardKey.Find)
        find_action.triggered.connect(self.open_find_dialog)
        self.addAction(find_action)

        # Status Bar interactions
        self.view_group.idClicked.connect(self.on_view_mode_changed)
        self.zoom_slider.valueChanged.connect(self.on_zoom_changed)
        self.zoom_out_btn.clicked.connect(lambda: self.step_zoom(-10))
        self.zoom_in_btn.clicked.connect(lambda: self.step_zoom(10))
        self.sync_toggle.toggled.connect(self.on_sync_toggled)

        # Ctrl + wheel zoom
        self.grid1.installEventFilter(self)
        self.grid2.installEventFilter(self)
        self.grid1.viewport().installEventFilter(self)
        self.grid2.viewport().installEventFilter(self)

    def show_sync_params_menu(self, pos):
        menu = QMenu(self)
        menu.setTitle("Synchronization Options")
        
        # Actions
        act_scroll = QAction("Sync Scrolling", self, checkable=True)
        act_scroll.setChecked(self.sync_scroll_enabled)
        act_scroll.toggled.connect(lambda v: setattr(self, 'sync_scroll_enabled', v))
        
        act_row = QAction("Sync Row Heights", self, checkable=True)
        act_row.setChecked(self.sync_row_height_enabled)
        act_row.toggled.connect(lambda v: setattr(self, 'sync_row_height_enabled', v))
        
        act_col = QAction("Sync Column Widths", self, checkable=True)
        act_col.setChecked(self.sync_col_width_enabled)
        act_col.toggled.connect(lambda v: setattr(self, 'sync_col_width_enabled', v))
        
        menu.addAction(act_scroll)
        menu.addAction(act_row)
        menu.addAction(act_col)
        
        menu.exec(self.sync_toggle.mapToGlobal(pos))

    def show_context_menu(self, pos, grid):
        menu = QMenu(self)
        
        copy_action = menu.addAction("Copy")
        copy_action.triggered.connect(self.copy_selected_range) # Uses active grid, which should be correct
        
        # Add more actions here in the future
        
        menu.exec(grid.mapToGlobal(pos))


    # ... (sync_selection, load_data remain same) ...
    def sync_selection(self, source, target):
        if self.is_syncing_selection:
            return
        
        self.is_syncing_selection = True
        target.clearSelection()
        for range_ in source.selectedRanges():
            target.setRangeSelected(range_, True)
        self.is_syncing_selection = False

    def on_grid_selection_changed(self, source, target):
        if self.is_syncing_selection:
            return
        self.active_grid = source
        self.sync_selection(source, target)
        self.update_formula_bar(source)
        if not self.sync_toggle.isChecked():
            self.sync_zoom_ui_for_grid(source)

    def update_formula_bar(self, source=None):
        grid = source or self._get_active_grid()
        if grid is None:
            self.name_box.setText("")
            self.formula_bar.setText("")
            return
        ranges = grid.selectedRanges()
        if not ranges:
            self.name_box.setText("")
            self.formula_bar.setText("")
            return
        r = ranges[0]
        self.name_box.setText(self.format_range(r))

        row = r.topRow()
        col = r.leftColumn()
        if grid is self.grid1:
            formula = self.formulas1.get((row, col))
            val = self.get_cell_value(self.df1, row, col)
        else:
            formula = self.formulas2.get((row, col))
            val = self.get_cell_value(self.df2, row, col)


        if formula and str(formula).startswith("="):
            self.highlight_formula_references(str(formula), grid)
        else:
            self.clear_reference_highlights()
            self.formula_bar.setText("" if val is None else str(val))

    def clear_reference_highlights(self, grid=None):
        target_grid = self.highlighted_grid or grid
        if not target_grid: return
        
        for r, c in self.highlighted_refs:
            item = target_grid.item(r, c)
            if item:
                item.setData(Qt.ItemDataRole.UserRole + 3, None)
        self.highlighted_refs = []
        self.highlighted_grid = None

    def highlight_formula_references(self, formula_text, grid):
        self.clear_reference_highlights()
        if not grid: 
            self.formula_bar.setText(formula_text)
            return
            
        self.highlighted_grid = grid

        # Regex to find cell references (e.g. A1, $A$1, A1:B2)
        # We process matches and rebuild HTML
        pattern = re.compile(r'(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)')
        
        parts = pattern.split(formula_text)
        matches = pattern.findall(formula_text)
        
        html_parts = []
        color_idx = 0
        
        # If split works correctly, it alternates non-match, match, non-match...
        # But 'split' with capturing group includes the matches in the list.
        # e.g. "=A1+B2" -> ['=', 'A1', '+', 'B2', '']
        
        for part in parts:
            if not part:
                continue
            
            if pattern.fullmatch(part):
                # This is a reference
                color = self.ref_colors[color_idx % len(self.ref_colors)]
                color_idx += 1
                
                # Highlight in Grid
                self.highlight_range_in_grid(grid, part, color)
                
                # Add to HTML
                html_parts.append(f'<span style="color:{color}; font-weight:bold;">{html.escape(part)}</span>')
            else:
                html_parts.append(html.escape(part))
        
        self.formula_bar.setHtml("".join(html_parts))

    def highlight_range_in_grid(self, grid, ref_str, color):
        # Handle cleanup relative to sheet? For now assume current sheet.
        try:
            # Remove $ for parsing
            clean_ref = ref_str.replace("$", "")
            
            if ":" in clean_ref:
                start, end = clean_ref.split(":")
                min_col, min_row = coordinate_from_string(start)
                max_col, max_row = coordinate_from_string(end)
                min_c = column_index_from_string(min_col) - 1
                max_c = column_index_from_string(max_col) - 1
                min_r = min_row - 1
                max_r = max_row - 1
            else:
                col_letter, row_num = coordinate_from_string(clean_ref)
                min_c = column_index_from_string(col_letter) - 1
                max_c = min_c
                min_r = row_num - 1
                max_r = min_r
            
            # Clamp to grid size
            rows = grid.rowCount()
            cols = grid.columnCount()
            
            for r in range(min_r, max_r + 1):
                if r < 0 or r >= rows: continue
                for c in range(min_c, max_c + 1):
                    if c < 0 or c >= cols: continue
                    
                    item = grid.item(r, c)
                    if item:
                        item.setData(Qt.ItemDataRole.UserRole + 3, color)
                        self.highlighted_refs.append((r, c))
                        
        except Exception as e:
            # If parsing fails (e.g. named range, or bad ref), just ignore
            print(f"Failed to highlight ref {ref_str}: {e}")
            pass

    def format_range(self, range_obj):
        start = f"{self.get_col_letter(range_obj.leftColumn() + 1)}{range_obj.topRow() + 1}"
        if range_obj.leftColumn() == range_obj.rightColumn() and range_obj.topRow() == range_obj.bottomRow():
            return start
        end = f"{self.get_col_letter(range_obj.rightColumn() + 1)}{range_obj.bottomRow() + 1}"
        return f"{start}:{end}"

    def load_data(self, df1, df2, styles1, styles2, merges1, merges2, formulas1, formulas2, col_widths1, row_heights1, col_widths2, row_heights2, comparison_result, file1_name=None, file2_name=None):
        # ... Same as before ...
        summary = comparison_result['summary']
        self.status_info.setText("Loading...")
        QGuiApplication.processEvents()

        self.apply_file_labels(file1_name, file2_name)
        
        total_rows = summary['total_rows']
        total_cols = summary['total_cols']

        changes = comparison_result['changes']
        
        change_set = {}
        for c in changes:
            change_set[(c['row'], c['col'])] = c

        self.df1 = df1
        self.df2 = df2
        self.styles1 = styles1
        self.styles2 = styles2
        self.merges1 = merges1
        self.merges2 = merges2
        self.formulas1 = formulas1 or {}
        self.formulas2 = formulas2 or {}
        self.change_set = change_set
        self.total_rows = total_rows
        self.total_cols = total_cols
        self.base_col_widths1 = col_widths1 or {}
        self.base_row_heights1 = row_heights1 or {}
        self.base_col_widths2 = col_widths2 or {}
        self.base_row_heights2 = row_heights2 or {}
        self.decimal_overrides1 = {}
        self.decimal_overrides2 = {}
        self.active_grid = None
        self.update_formula_bar(None)

        self.diff_map = self.build_diff_map(df1, df2, styles1, styles2, formulas1, formulas2, total_rows, total_cols)
        self.diff_rows_content, self.diff_rows_format, self.diff_rows_formula = self.build_diff_rows(self.diff_map)

        self.populate_grid(self.grid1, df1, styles1, merges1, change_set, total_rows, total_cols, self.decimal_overrides1, col_widths1, row_heights1, is_new=False)
        self.populate_grid(self.grid2, df2, styles2, merges2, change_set, total_rows, total_cols, self.decimal_overrides2, col_widths2, row_heights2, is_new=True)
        self.apply_row_filter()
        # Cache base sizes for zoom calculations
        self.base_table_font = QFont(self.grid1.font())
        self.base_default_row_height1 = self.grid1.verticalHeader().defaultSectionSize()
        self.base_default_row_height2 = self.grid2.verticalHeader().defaultSectionSize()
        self.base_col_widths1_full = [self.grid1.columnWidth(i) for i in range(total_cols)]
        self.base_col_widths2_full = [self.grid2.columnWidth(i) for i in range(total_cols)]
        self.base_row_heights1_px = {idx: self.grid1.rowHeight(idx) for idx in self.base_row_heights1.keys()}
        self.base_row_heights2_px = {idx: self.grid2.rowHeight(idx) for idx in self.base_row_heights2.keys()}
        self.update_status_info()
        self.set_zoom(100)

    def apply_file_labels(self, file1_name, file2_name):
        default_left = "Original"
        default_right = "Modified"

        self.file1_path = file1_name
        self.file2_path = file2_name

        name1 = os.path.basename(file1_name) if file1_name else default_left
        name2 = os.path.basename(file2_name) if file2_name else default_right

        stem1 = Path(name1).stem
        stem2 = Path(name2).stem

        ratio = 0.0
        if stem1 and stem2:
            ratio = SequenceMatcher(None, stem1.lower(), stem2.lower()).ratio()

        if ratio >= 0.75:
            self.grid1_label.setText(f"{name1} (old)")
            self.grid2_label.setText(f"{name2} (new)")
        else:
            self.grid1_label.setText(name1)
            self.grid2_label.setText(name2)

    def show_file_context_menu(self, side, pos):
        path = self.file1_path if side == 1 else self.file2_path
        label = self.grid1_label if side == 1 else self.grid2_label

        menu = QMenu(self)
        act_open = menu.addAction("Open File")
        act_open_loc = menu.addAction("Open Location")
        act_copy = menu.addAction("Copy Path")

        if not path:
            act_open.setEnabled(False)
            act_open_loc.setEnabled(False)
            act_copy.setEnabled(False)

        action = menu.exec(label.mapToGlobal(pos))
        if not action:
            return
        if not path:
            return

        if action == act_copy:
            QGuiApplication.clipboard().setText(path)
            return

        if not os.path.exists(path):
            QMessageBox.information(self, "File Not Found", "The file path is no longer available.")
            return

        if action == act_open:
            try:
                os.startfile(path)
            except Exception:
                QMessageBox.information(self, "Open File", "Unable to open the file.")
        elif action == act_open_loc:
            try:
                subprocess.run(["explorer", "/select,", os.path.normpath(path)], check=False)
            except Exception:
                try:
                    os.startfile(os.path.dirname(path))
                except Exception:
                    QMessageBox.information(self, "Open Location", "Unable to open the file location.")

    def set_sheet_options(self, side, sheets, current):
        combo = self.grid1_sheet_combo if side == 1 else self.grid2_sheet_combo
        btn = self.grid1_sheet_btn if side == 1 else self.grid2_sheet_btn
        combo.blockSignals(True)
        combo.clear()
        if sheets:
            combo.addItems(sheets)
            combo.setEnabled(True)
            if current in sheets:
                combo.setCurrentText(current)
            else:
                combo.setCurrentIndex(0)
        else:
            combo.addItem("No sheets found")
            combo.setEnabled(False)
        combo.blockSignals(False)
        current_name = combo.currentText() if combo.isEnabled() else None
        btn.setText(current_name or "No sheets")
        btn.setEnabled(bool(sheets))
        if side == 1:
            self.sheet1_current = current_name
        else:
            self.sheet2_current = current_name
        self.hide_sheet_combo(side)

    def on_sheet_combo_changed(self, side, name):
        if not name or name == "No sheets found":
            return
        btn = self.grid1_sheet_btn if side == 1 else self.grid2_sheet_btn
        btn.setText(name)
        if side == 1:
            if name == self.sheet1_current:
                return
            self.sheet1_current = name
        else:
            if name == self.sheet2_current:
                return
            self.sheet2_current = name
        self.sheetChangeRequested.emit(side, name)

    def on_sheet_combo_selected(self, side, idx):
        combo = self.grid1_sheet_combo if side == 1 else self.grid2_sheet_combo
        name = combo.itemText(idx)
        self.hide_sheet_combo(side)
        self.on_sheet_combo_changed(side, name)

    def show_sheet_combo(self, side):
        combo = self.grid1_sheet_combo if side == 1 else self.grid2_sheet_combo
        btn = self.grid1_sheet_btn if side == 1 else self.grid2_sheet_btn
        if not combo.isEnabled():
            return
        btn.setVisible(False)
        combo.setVisible(True)
        combo.setFocus()
        combo.showPopup()

    def hide_sheet_combo(self, side):
        combo = self.grid1_sheet_combo if side == 1 else self.grid2_sheet_combo
        btn = self.grid1_sheet_btn if side == 1 else self.grid2_sheet_btn
        combo.setVisible(False)
        btn.setVisible(True)

    def eventFilter(self, obj, event):
        if event.type() == QEvent.Type.Wheel:
            if event.modifiers() & Qt.KeyboardModifier.ControlModifier:
                table = None
                if obj in (self.grid1, self.grid1.viewport()):
                    table = self.grid1
                elif obj in (self.grid2, self.grid2.viewport()):
                    table = self.grid2
                if table is not None:
                    delta = event.angleDelta().y()
                    if delta:
                        step = 10 if delta > 0 else -10
                        if self.sync_toggle.isChecked():
                            self.step_zoom(step)
                        else:
                            self.step_zoom_for_table(table, step)
                        event.accept()
                        return True
        if not hasattr(self, "grid1_sheet_combo") or not hasattr(self, "grid2_sheet_combo"):
            return super().eventFilter(obj, event)
        if obj in (self.grid1_sheet_combo, self.grid2_sheet_combo):
            if event.type() == QEvent.Type.FocusOut:
                # If popup is visible, don't hide
                if obj.view() and obj.view().isVisible():
                    return False
                side = 1 if obj == self.grid1_sheet_combo else 2
                self.hide_sheet_combo(side)
            elif event.type() == QEvent.Type.Hide:
                 # Ensure button comes back if combo is hidden
                 side = 1 if obj == self.grid1_sheet_combo else 2
                 self.hide_sheet_combo(side)

        return super().eventFilter(obj, event)

    def on_view_mode_changed(self, mode_id):
        self.view_mode = "normal" if mode_id == 0 else "page_break"
        self.update_status_info()
        self.viewModeChanged.emit(self.view_mode)

    def set_view_mode(self, mode):
        mode = "page_break" if mode == "page_break" else "normal"
        if mode == self.view_mode:
            return
        self.view_mode = mode
        self.view_group.blockSignals(True)
        if mode == "normal":
            self.view_normal_btn.setChecked(True)
        else:
            self.view_page_btn.setChecked(True)
        self.view_group.blockSignals(False)
        self.update_status_info()
        self.viewModeChanged.emit(self.view_mode)

    def update_status_info(self):
        if not self.total_rows or not self.total_cols:
            self.status_info.setText("Ready")
            return

        visible_rows = self._get_visible_row_count()
        if visible_rows and visible_rows != self.total_rows:
            base = f"{visible_rows} of {self.total_rows} rows x {self.total_cols} cols"
        else:
            base = f"{self.total_rows} rows x {self.total_cols} cols"

        tags = []
        if self.only_changes:
            tags.append("Only changes")
        elif self.show_differences:
            tags.append("Diff")
        if self._has_active_column_filters():
            tags.append("Filtered")
        if getattr(self, "view_mode", "normal") == "page_break":
            tags.append("Page Break")

        if tags:
            base = f"{base}  •  " + " | ".join(tags)
        self.status_info.setText(base)

    def set_loading(self, active):
        if not hasattr(self, "status_progress"):
            return
        if active:
            self.status_progress.setRange(0, 0)  # indeterminate
            self.status_progress.setVisible(True)
        else:
            self.status_progress.setRange(0, 100)
            self.status_progress.setValue(0)
            self.status_progress.setVisible(False)

    def _has_active_column_filters(self):
        header = self.grid1.horizontalHeader()
        if isinstance(header, FilterHeader):
            return bool(header._active_filters)
        return False

    def _get_visible_row_count(self):
        if self.total_rows == 0:
            return 0
        if self.grid1.rowCount() == 0:
            return self.total_rows
        visible = 0
        for r in range(self.total_rows):
            if not self.grid1.isRowHidden(r):
                visible += 1
        return visible

    def on_zoom_changed(self, value):
        if self.sync_toggle.isChecked():
            self.set_zoom(value)
        else:
            table = self._get_active_grid() or self.grid1
            self.set_zoom_for_table(table, value, update_slider=False)

    def step_zoom(self, delta):
        if self.sync_toggle.isChecked():
            target = self.zoom_percent + delta
            target = self.clamp_zoom_percent(target)
            self.set_zoom(target)
        else:
            table = self._get_active_grid() or self.grid1
            self.step_zoom_for_table(table, delta)

    def step_zoom_for_table(self, table, delta):
        current = self.get_zoom_for_table(table)
        target = self.clamp_zoom_percent(current + delta)
        self.set_zoom_for_table(table, target)

    def set_zoom(self, percent):
        try:
            percent = int(percent)
        except Exception:
            return
        percent = self.clamp_zoom_percent(percent)
        if getattr(self, "zoom_percent", None) == percent and self.base_table_font is not None:
            self.update_zoom_ui(percent)
            return
        self.zoom_percent = percent
        self.zoom_percent1 = percent
        self.zoom_percent2 = percent
        self.update_zoom_ui(percent)
        scale = percent / 100.0
        self.apply_zoom_to_table(
            self.grid1,
            scale,
            getattr(self, "base_col_widths1_full", None),
            getattr(self, "base_default_row_height1", None),
            getattr(self, "base_row_heights1_px", None),
        )
        self.apply_zoom_to_table(
            self.grid2,
            scale,
            getattr(self, "base_col_widths2_full", None),
            getattr(self, "base_default_row_height2", None),
            getattr(self, "base_row_heights2_px", None),
        )

    def set_zoom_for_table(self, table, percent, update_slider=True):
        try:
            percent = int(percent)
        except Exception:
            return
        percent = self.clamp_zoom_percent(percent)
        if table is self.grid1:
            if getattr(self, "zoom_percent1", None) == percent and self.base_table_font is not None:
                self.update_zoom_ui(percent, update_slider=update_slider)
                return
            self.zoom_percent1 = percent
        else:
            if getattr(self, "zoom_percent2", None) == percent and self.base_table_font is not None:
                self.update_zoom_ui(percent, update_slider=update_slider)
                return
            self.zoom_percent2 = percent
        self.update_zoom_ui(percent, update_slider=update_slider)
        scale = percent / 100.0
        if table is self.grid1:
            self.apply_zoom_to_table(
                self.grid1,
                scale,
                getattr(self, "base_col_widths1_full", None),
                getattr(self, "base_default_row_height1", None),
                getattr(self, "base_row_heights1_px", None),
            )
        else:
            self.apply_zoom_to_table(
                self.grid2,
                scale,
                getattr(self, "base_col_widths2_full", None),
                getattr(self, "base_default_row_height2", None),
                getattr(self, "base_row_heights2_px", None),
            )

    def update_zoom_ui(self, percent, update_slider=True):
        self.zoom_label.setText(f"{percent}%")
        if update_slider:
            self.zoom_slider.blockSignals(True)
            self.zoom_slider.setValue(percent)
            self.zoom_slider.blockSignals(False)

    def clamp_zoom_percent(self, percent):
        if percent is None:
            return self.zoom_slider.value()
        return max(self.zoom_slider.minimum(), min(self.zoom_slider.maximum(), int(percent)))

    def get_zoom_for_table(self, table):
        return self.zoom_percent1 if table is self.grid1 else self.zoom_percent2

    def sync_zoom_ui_for_grid(self, table):
        if self.sync_toggle.isChecked():
            return
        percent = self.get_zoom_for_table(table)
        self.update_zoom_ui(percent)

    def on_sync_toggled(self, checked):
        if checked:
            active = self._get_active_grid() or self.grid1
            percent = self.get_zoom_for_table(active)
            self.set_zoom(percent)
        else:
            active = self._get_active_grid() or self.grid1
            self.sync_zoom_ui_for_grid(active)

    def apply_zoom_to_table(self, table, scale, base_col_widths, base_default_row_height, base_row_heights):
        if table is None:
            return
        if self.base_table_font is not None:
            font = QFont(self.base_table_font)
            # Safe resize: check pixel size first
            if font.pixelSize() > 0:
                new_pixel = max(8, int(font.pixelSize() * scale))
                font.setPixelSize(new_pixel)
            else:
                pts = font.pointSizeF()
                if pts <= 0:
                    pts = 11.0 # fallback default
                new_pts = max(6.0, pts * scale)
                font.setPointSizeF(new_pts)
            table.setFont(font)
            table.horizontalHeader().setFont(font)
            table.verticalHeader().setFont(font)

        if base_col_widths:
            for i, w in enumerate(base_col_widths):
                table.setColumnWidth(i, max(30, int(w * scale)))

        if base_default_row_height:
            table.verticalHeader().setDefaultSectionSize(max(18, int(base_default_row_height * scale)))

        if base_row_heights:
            for row_idx, height_px in base_row_heights.items():
                table.setRowHeight(row_idx, max(18, int(height_px * scale)))

    def populate_grid(self, table: QTableWidget, df, styles, merges, change_set, max_rows, max_cols, decimal_overrides, col_widths=None, row_heights=None, is_new=False):
        table.setUpdatesEnabled(False)
        table.setSortingEnabled(False)
        table.clearContents()
        table.setRowCount(max_rows)
        table.setColumnCount(max_cols)
        table.setShowGrid(False)
        
        table.setAlternatingRowColors(False)
        table.setStyleSheet("""
            QTableWidget {
                gridline-color: transparent;
                background-color: #ffffff;
                border: none;
                border-radius: 0px;
            }
            QTableWidget::item {
                padding: 0px;
            }
            QTableWidget::item:selected {
                background-color: #0ea5e9;
                color: #ffffff;
            }
            QHeaderView::section {
                background-color: #f1f5f9;
                padding: 0px;
                border: 1px solid #e2e8f0;
                font-weight: 600;
                color: #334155;
            }
        """)
        
        headers = [self.get_col_letter(i + 1) for i in range(max_cols)]
        # table.setHorizontalHeaderLabels(headers) # We set model via items usually, but HeaderView needs labels.
        # However, we are replacing the header, so we need to ensure labels persist.
        # QTableWidget stores labels in a model item.
        
        # Setup FilterHeader
        header = FilterHeader(Qt.Orientation.Horizontal, table)
        table.setHorizontalHeader(header)
        table.setHorizontalHeaderLabels(headers)
        for i, label in enumerate(headers):
            table.setHorizontalHeaderItem(i, QTableWidgetItem(label))
        header.setDefaultAlignment(Qt.AlignmentFlag.AlignCenter)
        table.horizontalHeader().setVisible(True)
        table.verticalHeader().setVisible(True)
        
        # Connect Headers for Sync (Resize) - we need to re-connect because we replaced the header object
        if hasattr(self, 'sync_col_resize'):
            # Determine if this is grid1 or grid2 based on is_new
            other_table = self.grid1 if is_new else self.grid2
            header.sectionResized.connect(lambda idx, old, new: self.sync_col_resize(table, other_table, idx, new))
            
        # Connect Filter Signal
        header.filterChanged.connect(lambda col, vals: self.on_column_filter_changed(table, col, vals))
        
        table.blockSignals(True)
        
        # Apply Custom Dimensions
        if col_widths:
            for col_idx, width in col_widths.items():
                if col_idx < max_cols:
                    # Approx 7px per Excel unit + padding? 
                    # Excel 8.43 ~ 64px. Ratio ~7.5
                    pixel_width = int(width * 7 + 5)
                    table.setColumnWidth(col_idx, pixel_width)
        else:
             # Initial default width
             pass
             
        if row_heights:
            for row_idx, height in row_heights.items():
                if row_idx < max_rows:
                    # Point to Pixel: 1 pt = 1.333 px
                    pixel_height = int(height * 1.33)
                    table.setRowHeight(row_idx, pixel_height)
        else:
             table.verticalHeader().setDefaultSectionSize(22)
            
        table.horizontalHeader().setFixedHeight(25)
        table.verticalHeader().setFixedWidth(30)
        
        df_rows, df_cols = df.shape

        for r in range(max_rows):
            for c in range(max_cols):
                cell_value = None
                if r < df_rows and c < df_cols:
                    v = df.iloc[r, c]
                    if str(v) not in ('nan', 'NaT', 'None'):
                         cell_value = v

                # Apply Styles
                bg_color = None
                cell_style = styles.get((r, c), {}) if styles else {}

                decimal_override = decimal_overrides.get((r, c), 0) if decimal_overrides else 0
                display_value = self.format_cell_value(cell_value, cell_style, decimal_override)
                item = QTableWidgetItem(display_value)
                
                # Apply Styles
                if cell_style:
                    
                    if 'bg' in cell_style:
                        bg_color = cell_style['bg']
                        item.setBackground(QColor(cell_style['bg']))
                    
                    if 'font' in cell_style:
                        f_data = cell_style['font']
                        font = item.font()
                        font_modified = False
                        if f_data.get('bold'):
                            font.setBold(True)
                            font_modified = True
                        if f_data.get('italic'):
                            font.setItalic(True)
                            font_modified = True
                        size_val = f_data.get('size')
                        if isinstance(size_val, (int, float)) and math.isfinite(size_val):
                            try:
                                size_pt = int(round(size_val))
                                if size_pt > 0:
                                    font.setPointSize(size_pt)
                                    font_modified = True
                            except Exception:
                                pass

                        if f_data.get('name'):
                            font.setFamily(f_data['name'])
                            font_modified = True
                        
                        if font_modified:
                            # If font became invalid (no point size and no pixel size), set a default
                            if font.pointSize() <= 0 and font.pixelSize() <= 0:
                                font.setPointSize(11)
                            item.setFont(font)
                        if f_data.get('color'):
                            fg = f_data['color']
                            if self.is_light_text_on_light_bg(fg, bg_color):
                                item.setForeground(QColor("#000000"))
                            else:
                                item.setForeground(QColor(fg))

                    if 'align' in cell_style:
                        align_data = cell_style['align']
                        qt_align = Qt.AlignmentFlag.AlignVCenter
                        v = align_data.get('vertical')
                        if v == 'top': qt_align = Qt.AlignmentFlag.AlignTop
                        elif v == 'bottom': qt_align = Qt.AlignmentFlag.AlignBottom
                        
                        h = align_data.get('horizontal')
                        if h == 'left': qt_align |= Qt.AlignmentFlag.AlignLeft
                        elif h == 'right': qt_align |= Qt.AlignmentFlag.AlignRight
                        elif h == 'center': qt_align |= Qt.AlignmentFlag.AlignHCenter
                        else: qt_align |= Qt.AlignmentFlag.AlignLeft
                        
                        item.setTextAlignment(qt_align)

                    # BORDERS (Pass to Delegate)
                    if 'border' in cell_style:
                        item.setData(Qt.ItemDataRole.UserRole, cell_style['border'])

                # Apply Diff Marker (Preserve original fill + text colors)
                diff_mask = self.get_cell_diff_mask(r, c)
                if diff_mask:
                    change_info = change_set.get((r, c))
                    item.setData(Qt.ItemDataRole.UserRole + 2, diff_mask)
                    tips = []
                    if change_info and (diff_mask & 1):
                        tips.append(f"Old: {change_info.get('old')}\nNew: {change_info.get('new')}")
                    if diff_mask & 4:
                        old_f = self.formulas1.get((r, c), "")
                        new_f = self.formulas2.get((r, c), "")
                        tips.append(f"Formula:\n{old_f}\n→ {new_f}")
                    if diff_mask & 2:
                        tips.append("Format changed")
                    if tips:
                        item.setToolTip("\n\n".join(tips))
                
                table.setItem(r, c, item)

        # Apply Merges
        # merges: list of (r, c, r_span, c_span)
        for (r, c, r_span, c_span) in merges:
            if r < max_rows and c < max_cols:
                table.setSpan(r, c, r_span, c_span)

        table.blockSignals(False)
        table.setUpdatesEnabled(True)
        
        # Removed self.fit_columns(table, max_cols) to respect Excel dimensions

    def get_col_letter(self, n):
        string = ""
        while n > 0:
            n, remainder = divmod(n - 1, 26)
            string = chr(65 + remainder) + string
        return string

    def adjust_decimals(self, delta):
        ranges = self.capture_selection()
        if not ranges:
            return
        self.apply_decimal_overrides(self.decimal_overrides1, ranges, delta)
        self.apply_decimal_overrides(self.decimal_overrides2, ranges, delta)
        self.refresh_view(ranges)

    def auto_fit_columns(self):
        self.fit_columns(self.grid1, self.total_cols)
        self.fit_columns(self.grid2, self.total_cols)

    def fit_columns(self, table, max_cols):
        table.resizeColumnsToContents()
        for i in range(max_cols):
            if table.columnWidth(i) < 50:
                table.setColumnWidth(i, 50)

    def set_show_differences(self, enabled):
        self.show_differences = bool(enabled)
        ranges = self.capture_selection()
        if not self.show_differences:
            self.only_changes = False
        self.refresh_view(ranges)
        self.update_status_info()

    def is_light_text_on_light_bg(self, fg, bg):
        if not fg:
            return False
        fg_val = fg.lower()
        if fg_val not in ("#ffffff", "ffffff"):
            return False
        if not bg:
            return True
        bg_val = bg.lower()
        if bg_val in ("#ffffff", "ffffff", "#f8fafc", "f8fafc", "#f1f5f9", "f1f5f9"):
            return True
        return False

    def set_diff_options(self, content=None, format=None, formula=None, only_changes=None, ignore_case=None, ignore_whitespace=None, ignore_number_format=None):
        if content is not None:
            self.show_diff_content = bool(content)
        if format is not None:
            self.show_diff_format = bool(format)
        if formula is not None:
            self.show_diff_formula = bool(formula)
        if only_changes is not None:
            self.only_changes = bool(only_changes)
        if ignore_case is not None:
            self.ignore_case = bool(ignore_case)
        if ignore_whitespace is not None:
            self.ignore_whitespace = bool(ignore_whitespace)
        if ignore_number_format is not None:
            self.ignore_number_format = bool(ignore_number_format)
        if not self.show_differences:
            self.only_changes = False
        # Rebuild diff map when ignore options change
        self.diff_map = self.build_diff_map(self.df1, self.df2, self.styles1, self.styles2, self.formulas1, self.formulas2, self.total_rows, self.total_cols)
        self.diff_rows_content, self.diff_rows_format, self.diff_rows_formula = self.build_diff_rows(self.diff_map)
        ranges = self.capture_selection()
        self.refresh_view(ranges)
        self.update_status_info()

    def get_cell_diff_mask(self, row, col):
        if not self.show_differences:
            return 0
        mask = self.diff_map.get((row, col), 0)
        if mask == 0:
            return 0
        if self.show_diff_content and self.show_diff_format:
            if self.show_diff_formula:
                return mask
            return mask & 0b011
        if self.show_diff_content:
            return mask & 0b001
        if self.show_diff_format:
            return mask & 0b010
        if self.show_diff_formula:
            return mask & 0b100
        return 0

    def apply_row_filter(self):
        if self.total_rows == 0:
            return
        if not self.only_changes:
            for row in range(self.total_rows):
                self.grid1.setRowHidden(row, False)
                self.grid2.setRowHidden(row, False)
            self.update_status_info()
            return
        allowed = set()
        if self.show_diff_content:
            allowed = allowed.union(self.diff_rows_content)
        if self.show_diff_format:
            allowed = allowed.union(self.diff_rows_format)
        if self.show_diff_formula:
            allowed = allowed.union(self.diff_rows_formula)
        for row in range(self.total_rows):
            hide = row not in allowed
            self.grid1.setRowHidden(row, hide)
            self.grid2.setRowHidden(row, hide)
        self.update_status_info()

    def build_diff_rows(self, diff_map):
        rows_content = set()
        rows_format = set()
        rows_formula = set()
        for (r, _c), mask in diff_map.items():
            if mask & 1:
                rows_content.add(r)
            if mask & 2:
                rows_format.add(r)
            if mask & 4:
                rows_formula.add(r)
        return rows_content, rows_format, rows_formula

    def build_diff_map(self, df1, df2, styles1, styles2, formulas1, formulas2, max_rows, max_cols):
        diff_map = {}
        for r in range(max_rows):
            for c in range(max_cols):
                v1 = self.get_cell_value(df1, r, c)
                v2 = self.get_cell_value(df2, r, c)
                content_diff = not self.values_equal(v1, v2)

                f1 = (formulas1 or {}).get((r, c))
                f2 = (formulas2 or {}).get((r, c))
                formula_diff = not self.formulas_equal(f1, f2)

                s1 = styles1.get((r, c), {}) if styles1 else {}
                s2 = styles2.get((r, c), {}) if styles2 else {}
                format_diff = self.style_signature(s1) != self.style_signature(s2)

                mask = 0
                if content_diff:
                    mask |= 1
                if format_diff:
                    mask |= 2
                if formula_diff:
                    mask |= 4
                if mask:
                    diff_map[(r, c)] = mask
        return diff_map

    def get_cell_value(self, df, r, c):
        if df is None:
            return None
        if r >= df.shape[0] or c >= df.shape[1]:
            return None
        v = df.iloc[r, c]
        if str(v) in ('nan', 'NaT', 'None'):
            return None
        return v

    def values_equal(self, a, b):
        if a is None and b is None:
            return True
        if a is None or b is None:
            return False
        if isinstance(a, numbers.Number) and isinstance(b, numbers.Number):
            return float(a) == float(b)
        return self.normalize_text(str(a)) == self.normalize_text(str(b))

    def formulas_equal(self, a, b):
        if a is None and b is None:
            return True
        if a is None or b is None:
            return False
        return self.normalize_text(str(a)) == self.normalize_text(str(b))

    def normalize_text(self, value):
        if not isinstance(value, str):
            return value
        text = value
        if self.ignore_whitespace:
            text = " ".join(text.split())
        if self.ignore_case:
            text = text.lower()
        return text

    def style_signature(self, style):
        if not style:
            style = {}
        font = style.get('font', {})
        align = style.get('align', {})
        border = style.get('border', {})
        numfmt = None if self.ignore_number_format else style.get('numfmt')
        return (
            style.get('bg'),
            font.get('color'), font.get('bold'), font.get('italic'), font.get('size'), font.get('name'),
            align.get('horizontal'), align.get('vertical'), align.get('wrap_text'),
            bool(border.get('left')), bool(border.get('right')), bool(border.get('top')), bool(border.get('bottom')),
            numfmt,
        )

    def copy_selected_range(self):
        table = self._get_active_grid()
        if table is None:
            return
        ranges = table.selectedRanges()
        if not ranges:
            return
        r = ranges[0]
        lines = []
        for row in range(r.topRow(), r.bottomRow() + 1):
            cells = []
            for col in range(r.leftColumn(), r.rightColumn() + 1):
                item = table.item(row, col)
                cells.append(item.text() if item else "")
            lines.append("\t".join(cells))
        QGuiApplication.clipboard().setText("\n".join(lines))

    def _get_active_grid(self):
        if self.grid1.hasFocus() or self.grid1.selectedRanges():
            return self.grid1
        if self.grid2.hasFocus() or self.grid2.selectedRanges():
            return self.grid2
        return None

    def apply_decimal_overrides(self, overrides, ranges, delta):
        if overrides is None:
            return
        for top, left, bottom, right in ranges:
            for row in range(top, bottom + 1):
                for col in range(left, right + 1):
                    key = (row, col)
                    overrides[key] = overrides.get(key, 0) + delta
                    if overrides[key] == 0:
                        del overrides[key]

    def refresh_view(self, ranges=None):
        if self.df1 is None or self.df2 is None:
            return
        self.populate_grid(self.grid1, self.df1, self.styles1, self.merges1, self.change_set, self.total_rows, self.total_cols, self.decimal_overrides1, is_new=False)
        self.populate_grid(self.grid2, self.df2, self.styles2, self.merges2, self.change_set, self.total_rows, self.total_cols, self.decimal_overrides2, is_new=True)
        self.apply_row_filter()
        if ranges:
            self.restore_selection(ranges)

    def format_cell_value(self, value, style, decimal_override=0):
        if value is None:
            return ""
        if isinstance(value, str):
            return value
        if isinstance(value, bool):
            return str(value)

        if isinstance(value, numbers.Number):
            if decimal_override == 0:
                return str(value)

            numfmt = style.get('numfmt') if style else None
            if numfmt and self.is_date_format(numfmt):
                return str(value)
            fmt_info = self.parse_number_format(numfmt)

            decimals = fmt_info.get('decimals', 0) if fmt_info else 0
            decimals = max(0, decimals + decimal_override)
            grouping = fmt_info.get('grouping', False) if fmt_info else False
            prefix = fmt_info.get('prefix', "") if fmt_info else ""
            suffix = fmt_info.get('suffix', "") if fmt_info else ""
            prefix, suffix = self.strip_parentheses(prefix, suffix)
            is_percent = fmt_info.get('percent', False) if fmt_info else False

            val = float(value)
            if is_percent:
                val *= 100
                if "%" not in suffix:
                    suffix = f"{suffix}%"

            fmt = f"{{:,.{decimals}f}}" if grouping else f"{{:.{decimals}f}}"
            return f"{prefix}{fmt.format(val)}{suffix}"

        return str(value)

    def capture_selection(self):
        table = self._get_active_grid()
        if table is None:
            return []
        self.active_grid = table
        ranges = table.selectedRanges()
        return [(r.topRow(), r.leftColumn(), r.bottomRow(), r.rightColumn()) for r in ranges]

    def restore_selection(self, ranges):
        self.is_syncing_selection = True
        for table in (self.grid1, self.grid2):
            table.blockSignals(True)
            table.clearSelection()
            for top, left, bottom, right in ranges:
                selection = QTableWidgetSelectionRange(top, left, bottom, right)
                table.setRangeSelected(selection, True)
            table.blockSignals(False)
        self.is_syncing_selection = False
        self.update_formula_bar(self.active_grid)

    def strip_parentheses(self, prefix, suffix):
        if prefix:
            prefix = prefix.replace("(", "").replace(")", "")
        if suffix:
            suffix = suffix.replace("(", "").replace(")", "")
        return prefix, suffix

    def is_date_format(self, fmt):
        if not fmt:
            return False
        section = str(fmt).split(";")[0]
        section = re.sub(r'\".*?\"', '', section)
        section = re.sub(r'\[[^\]]+\]', '', section)
        return bool(re.search(r'[ymdhsa]', section, re.IGNORECASE))

    def parse_number_format(self, fmt):
        if not fmt:
            return None
        fmt = str(fmt)
        if fmt.lower() in ("general", "@"):
            return None

        section = fmt.split(";")[0]
        section = re.sub(r'\".*?\"', '', section)
        section = re.sub(r'\[[^\]]+\]', '', section)

        placeholders = list(re.finditer(r'[0#?]', section))
        if not placeholders:
            return None
        first = placeholders[0].start()
        last = placeholders[-1].end()
        prefix = section[:first].strip().replace("_", "").replace("*", "").replace("\\", "")
        suffix = section[last:].strip().replace("_", "").replace("*", "").replace("\\", "")

        decimals = 0
        dec_match = re.search(r'\.(?P<dec>[0#?]+)', section)
        if dec_match:
            decimals = len(dec_match.group('dec'))

        grouping = "," in section
        percent = "%" in section

        if not prefix:
            currency_match = re.search(r'[$€£¥₫₩₹]', fmt)
            if currency_match:
                prefix = currency_match.group(0)

        return {
            "decimals": decimals,
            "grouping": grouping,
            "percent": percent,
            "prefix": prefix,
            "suffix": suffix,
        }

    def open_find_dialog(self):
        if not self.find_dialog:
            self.find_dialog = FindDialog(self)
            self.find_dialog.findNext.connect(self.on_find_next)
            self.find_dialog.findAll.connect(self.on_find_all)
        self.find_dialog.show()
        self.find_dialog.raise_()
        self.find_dialog.activateWindow()

    def navigate_to_find_result(self, data):
        # data = (grid_index, r, c)
        grid_idx, r, c = data
        target = self.grid1 if grid_idx == 1 else self.grid2
        target.setFocus()
        target.setCurrentCell(r, c)
        
        # Ensure visible
        target.scrollToItem(target.item(r, c))

    def on_find_next(self, options):
        # Determine start grid and start pos
        active_grid = self._get_active_grid() or self.grid1
        start_row = active_grid.currentRow()
        start_col = active_grid.currentColumn()
        if start_row < 0: start_row = 0
        if start_col < 0: start_col = 0
        
        found = self.search_grid(active_grid, options, start_row, start_col, find_next=True)
        
        if not found:
             # Confirm wrap (or continue from beginning)
             # "We couldn't find what you were looking for. Click Find Next to continue searching from the beginning of the sheet?"
             # To mimic Excel exactly, if we started in middle, we wrap.
             # If we fail to find anything at all (search entire grid), show "We couldn't find..."
             
             # Let's simple check if we can find anything at all first? 
             # No, that's expensive.
             
             # Just try wrapped search immediately
             found_wrapped = self.search_grid(active_grid, options, 0, 0, find_next=True)
             
             if found_wrapped:
                 # It exists somewhere else. Ask user to wrap.
                 reply = QMessageBox.question(
                     self, "Find", 
                     "We couldn't find what you were looking for. Click Yes to continue searching from the beginning of the sheet?",
                     QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
                 )
                 
                 if reply == QMessageBox.StandardButton.Yes:
                     found = found_wrapped
             else:
                 # Not found anywhere
                 QMessageBox.information(self, "Find", f"We couldn't find what you were looking for.")

        if found:
            # (r, c)
            r, c = found
            active_grid.setCurrentCell(r, c)
            active_grid.setFocus()
            
            # Highlight this result in the Find Dialog list if it's open
            if self.find_dialog and self.find_dialog.results_table.isVisible():
                self.sync_find_dialog_selection(active_grid, r, c)

    def sync_find_dialog_selection(self, grid, r, c):
        table = self.find_dialog.results_table
        grid_idx = 1 if grid is self.grid1 else 2
        
        for row in range(table.rowCount()):
            data = table.item(row, 0).data(Qt.ItemDataRole.UserRole)
            # data = (grid_idx, r, c)
            if data == (grid_idx, r, c):
                table.selectRow(row)
                table.scrollToItem(table.item(row, 0))
                break

    def on_find_all(self, options):
        results = []
        
        # Search Grid 1
        matches1 = self.search_grid(self.grid1, options, 0, 0, find_next=False)
        label1 = self.grid1_label.text()
        for r, c in matches1:
            val = self.get_searchable_value(self.grid1, r, c, options)
            ref = self.format_range(QTableWidgetSelectionRange(r, c, r, c))
            results.append((ref, val, label1, (1, r, c)))

        # Search Grid 2
        matches2 = self.search_grid(self.grid2, options, 0, 0, find_next=False)
        label2 = self.grid2_label.text()
        for r, c in matches2:
            val = self.get_searchable_value(self.grid2, r, c, options)
            ref = self.format_range(QTableWidgetSelectionRange(r, c, r, c))
            results.append((ref, val, label2, (2, r, c)))
            
        # Update Dialog
        self.find_dialog.results_table.setRowCount(0)
        self.find_dialog.results_table.show()
        for ref, val, label, user_data in results:
            self.find_dialog.add_result_row(ref, val, label, user_data)
        self.find_dialog.adjustSize()

    def search_grid(self, grid, options, start_r, start_c, find_next=True):
        # Logic to iterate grid (By Row / Col)
        rows = grid.rowCount()
        cols = grid.columnCount()
        
        text = options['text']
        match_case = options['match_case']
        match_entire = options['match_entire']
        look_in = options['look_in'] # Values / Formulas
        search_by = options['search_by'] # By Rows / By Columns
        
        if not text: return None if find_next else []
        
        matches = []
        
        # Generates coordinates based on order
        coords = []
        if search_by == "By Rows":
            for r in range(rows):
                for c in range(cols):
                    coords.append((r, c))
        else:
            for c in range(cols):
                for r in range(rows):
                    coords.append((r, c))
                    
        # If finding next, reorder to start from start_r, start_c
        if find_next:
            # Find index
            start_idx = 0
            # This is slow for large grids, optimization needed for production
            # But okay for prototype
            try:
                # Find simplistic index?
                 # Custom comparison
                 pass
            except:
                pass
            
            # Simple iteration from start position logic
            # We just iterate all coords, check if > start, else wrap
            # A bit complex to sort perfectly given By Row/Col options
            # Let's just do a linear scan from (0,0) and filtering > current if find_next
            # And handling wrap around manually in on_find_next
            
            # Re-implementing linear scan starting from current
            found_wrapped = None
            
            # Identify current index in coords
            # This handles the "Next" logic correctly respecting the sort order
            current_idx = -1
            for i, (r, c) in enumerate(coords):
                if r == start_r and c == start_c:
                    current_idx = i
                    break
            
            # Rotate list to start from next item
            if current_idx != -1:
                coords = coords[current_idx+1:] + coords[:current_idx+1]
            
            for r, c in coords:
                 match = self.check_match(grid, r, c, text, match_case, match_entire, look_in)
                 if match:
                     return (r, c)
            return None
            
        else:
            # Find All
            for r, c in coords:
                 if self.check_match(grid, r, c, text, match_case, match_entire, look_in):
                     matches.append((r, c))
            return matches

    def check_match(self, grid, r, c, text, match_case, match_entire, look_in):
        val = self.get_searchable_value(grid, r, c, {'look_in': look_in})
        
        if val is None:
            return False
        
        val_str = str(val)
        
        if not match_case:
            val_str = val_str.lower()
            text = text.lower()
            
        if match_entire:
            return val_str == text
        else:
            return text in val_str

    def get_searchable_value(self, grid, r, c, options):
        look_in = options['look_in']
        
        if look_in == "Formulas":
            # Check formula dicts
            if grid is self.grid1:
                f = self.formulas1.get((r, c))
            else:
                f = self.formulas2.get((r, c))
            
            if f and str(f).startswith("="):
                return f
            # If no formula, does it fallback to value? Excel usually searches values if no formula? 
            # Actually Excel searches the displayed formula string usually.
            # If no formula, use value
            return self.get_cell_value(self.df1 if grid is self.grid1 else self.df2, r, c)
            
        else: # Values
             return self.get_cell_value(self.df1 if grid is self.grid1 else self.df2, r, c)

    def toggle_freeze(self):
        # Implementation of Freeze Panes
        # For QTableWidget, "freezing" arbitrary panes is complex requiring split views.
        # For now, we will notify the user or implemented a basic header freeze if possible.
        # Since Excel-like arbitrary freeze is requested:
        
        QMessageBox.information(
            self, "Freeze Panes", 
            "Freeze Panes at Selection is not yet fully implemented.\n\n"
            "In a full version, this would split the view into quadrants based on your current selection ("
            f"{self.grid1.currentRow()+1}, {self.grid1.currentColumn()+1})."
        )

    def sync_scroll(self, source, target, orientation, value):
        if not self.sync_toggle.isChecked() or not self.sync_scroll_enabled:
            return
            
        if orientation == Qt.Orientation.Vertical:
            target.verticalScrollBar().setValue(value)
        else:
            target.horizontalScrollBar().setValue(value)

    def sync_col_resize(self, source, target, idx, new_size):
        if not self.sync_toggle.isChecked() or not self.sync_col_width_enabled:
            return
        if self.is_syncing_resize:
            return
        self.is_syncing_resize = True
        target.setColumnWidth(idx, new_size)
        self.is_syncing_resize = False

    def sync_row_resize(self, source, target, idx, new_size):
        if not self.sync_toggle.isChecked() or not self.sync_row_height_enabled:
            return
        if self.is_syncing_resize:
            return
        self.is_syncing_resize = True
        target.setRowHeight(idx, new_size)
        self.is_syncing_resize = False

    def on_column_filter_changed(self, table, col_idx, allowed_values):
        # table: which table triggered it
        # allowed_values: set of strings, or None if cleared
        
        # We need to maintain a state of filters for the table
        # Since FilterHeader is replaced on load, the state is in Header.
        # But we also need to combine filters across columns if multiple are active.
        # Actually FilterHeader maintains active_filters for Icon state.
        # But here we need to iterate rows and check against ALL active filters on that table.
        
        header = table.horizontalHeader()
        if not isinstance(header, FilterHeader):
            return
            
        all_filters = header._active_filters
        
        target_tables = [table]
        # Check Sync for Filtering
        # If Sync is ON, we sync row hiding to the other table as well
        if self.sync_toggle.isChecked() and self.sync_row_height_enabled:
             # Identify the other table
             other = self.grid2 if table == self.grid1 else self.grid1
             target_tables.append(other)
        
        # Loop over rows and Determine Visibility
        # Note: This is purely based on the Source Table's values.
        # If we sync, we apply the visibility mask of Source to Target.
        
        row_count = table.rowCount()
        
        table.setUpdatesEnabled(False)
        for t in target_tables:
            if t != table:
                t.setUpdatesEnabled(False)
        
        for r in range(row_count):
            should_show = True
            
            # Check all column filters
            for f_col, f_vals in all_filters.items():
                item = table.item(r, f_col)
                val = item.text() if item else ""
                
                if val not in f_vals:
                    should_show = False
                    break
            
            # Apply Visibility
            table.setRowHidden(r, not should_show)
            if self.sync_toggle.isChecked() and self.sync_row_height_enabled:
                 for t in target_tables:
                     if t != table:
                         t.setRowHidden(r, not should_show)

        table.setUpdatesEnabled(True)
        for t in target_tables:
            if t != table:
                t.setUpdatesEnabled(True)
        self.update_status_info()

    def request_swap(self):
        self.swapRequested.emit()
